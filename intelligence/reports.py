"""
Report Engine for Expo Design AI (Phase 8).

Generates professional engineering reports:
- PDF reports (reportlab)
- Excel reports (openpyxl)
- Word reports (python-docx)
- Compliance reports
- BIM QA reports
- Drawing review reports
- Issue registers

Reports include evidence and source references.
"""
import os
import json
import datetime
from typing import List, Dict, Any, Optional
from dataclasses import dataclass, field

try:
    from loguru import logger
except Exception:
    class _Nop:
        def __getattr__(self, _): return lambda *a, **k: None
    logger = _Nop()


@dataclass
class ReportSection:
    """A section in a report."""
    title: str
    content: str = ""
    data: Dict[str, Any] = field(default_factory=dict)
    table: List[Dict[str, Any]] = field(default_factory=list)


@dataclass
class Report:
    """A generated report."""
    report_id: str = ""
    title: str = ""
    report_type: str = ""
    project_id: str = ""
    discipline: str = ""
    sections: List[ReportSection] = field(default_factory=list)
    findings: List[Dict[str, Any]] = field(default_factory=list)
    metadata: Dict[str, Any] = field(default_factory=dict)
    created_at: str = ""
    created_by: str = ""
    format: str = "pdf"


class ReportEngine:
    """
    Report generation engine.
    Supports PDF, Excel, and Word output.
    """

    def __init__(self):
        self._templates_dir = os.path.join(os.path.dirname(__file__), "templates")
        os.makedirs(self._templates_dir, exist_ok=True)

    def generate_compliance_report(self, project_id: str,
                                   findings: List[Dict[str, Any]],
                                   discipline: str = "",
                                   created_by: str = "system") -> Report:
        """Generate a compliance report."""
        report = Report(
            report_id=f"RPT-{datetime.datetime.utcnow().strftime('%Y%m%d%H%M%S')}",
            title=f"Compliance Report — {project_id}",
            report_type="compliance",
            project_id=project_id,
            discipline=discipline,
            created_at=datetime.datetime.utcnow().isoformat(),
            created_by=created_by,
        )

        # Executive summary
        passed = sum(1 for f in findings if f.get("status") == "PASS")
        failed = sum(1 for f in findings if f.get("status") == "FAIL")
        review = sum(1 for f in findings if f.get("status") == "REVIEW")

        report.sections.append(ReportSection(
            title="Executive Summary",
            content=f"Total checks: {len(findings)}\nPassed: {passed}\nFailed: {failed}\nReview required: {review}",
            data={"total": len(findings), "passed": passed, "failed": failed, "review": review},
        ))

        # Findings detail
        for f in findings:
            report.sections.append(ReportSection(
                title=f.get("title", "Finding"),
                content=f.get("description", ""),
                data=f,
            ))

        # Recommendations
        recommendations = [f.get("recommendation", "") for f in findings if f.get("recommendation")]
        if recommendations:
            report.sections.append(ReportSection(
                title="Recommendations",
                content="\n".join(f"- {r}" for r in recommendations),
            ))

        report.findings = findings
        return report

    def generate_bim_qa_report(self, project_id: str,
                                bim_stats: Dict[str, Any],
                                findings: List[Dict[str, Any]],
                                created_by: str = "system") -> Report:
        """Generate a BIM QA report."""
        report = Report(
            report_id=f"RPT-{datetime.datetime.utcnow().strftime('%Y%m%d%H%M%S')}",
            title=f"BIM QA Report — {project_id}",
            report_type="bim_qa",
            project_id=project_id,
            created_at=datetime.datetime.utcnow().isoformat(),
            created_by=created_by,
        )

        # Model statistics
        report.sections.append(ReportSection(
            title="Model Statistics",
            content=json.dumps(bim_stats, indent=2),
            data=bim_stats,
        ))

        # Findings
        for f in findings:
            report.sections.append(ReportSection(
                title=f.get("title", "Finding"),
                content=f.get("description", ""),
                data=f,
            ))

        report.findings = findings
        return report

    def generate_drawing_review_report(self, project_id: str,
                                       drawings: List[Dict[str, Any]],
                                       findings: List[Dict[str, Any]],
                                       created_by: str = "system") -> Report:
        """Generate a drawing review report."""
        report = Report(
            report_id=f"RPT-{datetime.datetime.utcnow().strftime('%Y%m%d%H%M%S')}",
            title=f"Drawing Review Report — {project_id}",
            report_type="drawing_review",
            project_id=project_id,
            created_at=datetime.datetime.utcnow().isoformat(),
            created_by=created_by,
        )

        # Drawings reviewed
        report.sections.append(ReportSection(
            title="Drawings Reviewed",
            content=f"Total drawings: {len(drawings)}",
            table=drawings,
        ))

        # Findings
        for f in findings:
            report.sections.append(ReportSection(
                title=f.get("title", "Finding"),
                content=f.get("description", ""),
                data=f,
            ))

        report.findings = findings
        return report

    def generate_issue_register(self, project_id: str,
                                findings: List[Dict[str, Any]],
                                created_by: str = "system") -> Report:
        """Generate an issue register."""
        report = Report(
            report_id=f"RPT-{datetime.datetime.utcnow().strftime('%Y%m%d%H%M%S')}",
            title=f"Issue Register — {project_id}",
            report_type="issue_register",
            project_id=project_id,
            created_at=datetime.datetime.utcnow().isoformat(),
            created_by=created_by,
        )

        # Issue table
        table = []
        for f in findings:
            table.append({
                "ID": f.get("finding_id", ""),
                "Title": f.get("title", ""),
                "Status": f.get("status", ""),
                "Severity": f.get("severity", ""),
                "Source": f.get("source_doc", ""),
                "Page": f.get("source_page", ""),
            })

        report.sections.append(ReportSection(
            title="Issue Register",
            content=f"Total issues: {len(findings)}",
            table=table,
        ))

        report.findings = findings
        return report

    def export_pdf(self, report: Report, output_path: str) -> str:
        """Export report to PDF using reportlab."""
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib import colors
            from reportlab.lib.units import inch

            doc = SimpleDocTemplate(output_path, pagesize=A4,
                                    rightMargin=72, leftMargin=72,
                                    topMargin=72, bottomMargin=18)
            styles = getSampleStyleSheet()
            story = []

            # Title
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                spaceAfter=30,
            )
            story.append(Paragraph(report.title, title_style))
            story.append(Spacer(1, 12))

            # Metadata
            story.append(Paragraph(f"Project: {report.project_id}", styles['Normal']))
            story.append(Paragraph(f"Type: {report.report_type}", styles['Normal']))
            story.append(Paragraph(f"Created: {report.created_at}", styles['Normal']))
            story.append(Paragraph(f"By: {report.created_by}", styles['Normal']))
            story.append(Spacer(1, 24))

            # Sections
            for section in report.sections:
                story.append(Paragraph(section.title, styles['Heading2']))
                story.append(Spacer(1, 12))

                if section.content:
                    story.append(Paragraph(section.content, styles['Normal']))
                    story.append(Spacer(1, 12))

                if section.table:
                    # Create table
                    if section.table:
                        headers = list(section.table[0].keys()) if section.table else []
                        data = [headers]
                        for row in section.table:
                            data.append([str(row.get(h, "")) for h in headers])

                        table = Table(data)
                        table.setStyle(TableStyle([
                            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                            ('FONTSIZE', (0, 0), (-1, 0), 14),
                            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                            ('GRID', (0, 0), (-1, -1), 1, colors.black),
                        ]))
                        story.append(table)
                        story.append(Spacer(1, 12))

            doc.build(story)
            logger.info("PDF report generated: {}", output_path)
            return output_path

        except Exception as e:
            logger.error("PDF generation failed: {}", e)
            raise

    def export_excel(self, report: Report, output_path: str) -> str:
        """Export report to Excel using openpyxl."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill

            wb = Workbook()
            ws = wb.active
            ws.title = report.title[:31]  # Excel sheet name max 31 chars

            # Title
            ws['A1'] = report.title
            ws['A1'].font = Font(size=16, bold=True)
            ws.merge_cells('A1:F1')

            # Metadata
            ws['A3'] = f"Project: {report.project_id}"
            ws['A4'] = f"Type: {report.report_type}"
            ws['A5'] = f"Created: {report.created_at}"
            ws['A6'] = f"By: {report.created_by}"

            # Sections
            row = 8
            for section in report.sections:
                ws.cell(row=row, column=1, value=section.title).font = Font(size=14, bold=True)
                row += 1

                if section.content:
                    ws.cell(row=row, column=1, value=section.content)
                    row += 1

                if section.table:
                    headers = list(section.table[0].keys()) if section.table else []
                    for col, header in enumerate(headers, 1):
                        cell = ws.cell(row=row, column=col, value=header)
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
                    row += 1

                    for table_row in section.table:
                        for col, header in enumerate(headers, 1):
                            ws.cell(row=row, column=col, value=str(table_row.get(header, "")))
                        row += 1

                row += 2

            wb.save(output_path)
            logger.info("Excel report generated: {}", output_path)
            return output_path

        except Exception as e:
            logger.error("Excel generation failed: {}", e)
            raise

    def export_word(self, report: Report, output_path: str) -> str:
        """Export report to Word using python-docx."""
        try:
            from docx import Document
            from docx.shared import Inches, Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()

            # Title
            title = doc.add_heading(report.title, 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Metadata
            doc.add_paragraph(f"Project: {report.project_id}")
            doc.add_paragraph(f"Type: {report.report_type}")
            doc.add_paragraph(f"Created: {report.created_at}")
            doc.add_paragraph(f"By: {report.created_by}")
            doc.add_paragraph()

            # Sections
            for section in report.sections:
                doc.add_heading(section.title, level=1)

                if section.content:
                    doc.add_paragraph(section.content)

                if section.table:
                    headers = list(section.table[0].keys()) if section.table else []
                    table = doc.add_table(rows=1, cols=len(headers))
                    table.style = 'Light Grid Accent 1'

                    # Header row
                    hdr_cells = table.rows[0].cells
                    for i, header in enumerate(headers):
                        hdr_cells[i].text = header

                    # Data rows
                    for row_data in section.table:
                        row_cells = table.add_row().cells
                        for i, header in enumerate(headers):
                            row_cells[i].text = str(row_data.get(header, ""))

                doc.add_paragraph()

            doc.save(output_path)
            logger.info("Word report generated: {}", output_path)
            return output_path

        except Exception as e:
            logger.error("Word generation failed: {}", e)
            raise

    def export(self, report: Report, output_path: str, format: str = "pdf") -> str:
        """Export report in the specified format."""
        if format == "pdf":
            return self.export_pdf(report, output_path)
        elif format == "excel" or format == "xlsx":
            return self.export_excel(report, output_path)
        elif format == "word" or format == "docx":
            return self.export_word(report, output_path)
        else:
            raise ValueError(f"Unsupported format: {format}")


# Singleton instance
_report_engine = ReportEngine()


def get_report_engine() -> ReportEngine:
    """Get the global report engine instance."""
    return _report_engine
