"""
Excel Schedule Parser for Masterplan Processing Division (MPD).
Parses Masterplan Development Schedules, Plot Registers, and Zoning Matrices.
"""
import os
import re
from typing import List, Dict, Any, Optional
from loguru import logger
import openpyxl

from services.masterplan.models import MasterplanPlot


def _clean_num(val: Any, default: float = 0.0) -> float:
    if val is None:
        return default
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace(",", "").replace("sqm", "").replace("m2", "").replace("m²", "").strip()
    match = re.search(r"[-+]?\d*\.?\d+", s)
    if match:
        try:
            return float(match.group(0))
        except Exception:
            pass
    return default


class MasterplanExcelParser:
    """
    Ingests Excel plot registers and normalizes attributes.
    """

    def parse_excel(self, file_path: str, project_id: str) -> List[MasterplanPlot]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Excel file not found: {file_path}")

        wb = openpyxl.load_workbook(file_path, data_only=True)
        plots: List[MasterplanPlot] = []

        # Score all sheets to automatically find the one containing actual plot schedule data
        best_sheet = wb.active
        best_score = -1
        best_header_idx = -1
        best_col_map: Dict[str, int] = {}

        for name in wb.sheetnames:
            sh = wb[name]
            sh_rows = list(sh.iter_rows(values_only=True))
            if not sh_rows:
                continue

            h_idx = -1
            c_map: Dict[str, int] = {}
            for r_i, row in enumerate(sh_rows[:25]):
                row_strs = [str(c or "").strip().upper() for c in row]
                has_plot = any(k in c for c in row_strs for k in ("PLOT NO", "PLOT_ID", "PLOT ID", "PLOT #", "PARCEL", "PLOT"))
                has_area = any(k in c for c in row_strs for k in ("AREA", "GFA", "FAR", "BUA"))
                if has_plot and has_area:
                    h_idx = r_i
                    for c_i, c_val in enumerate(row_strs):
                        if not c_val:
                            continue
                        if any(k in c_val for k in ("PLOT NO", "PLOT_ID", "PLOT ID", "PLOT #", "PARCEL")):
                            c_map["plot_id"] = c_i
                        elif any(k in c_val for k in ("ZONE", "SECTOR", "DISTRICT", "PHASE")):
                            c_map["sector"] = c_i
                        elif any(k in c_val for k in ("LAND USE", "ZONING", "LANDUSE", "USE")):
                            c_map["land_use"] = c_i
                        elif any(k in c_val for k in ("LAND AREA", "PLOT AREA", "SITE AREA", "AREA (SQM)", "AREA SQM", "AREA")):
                            if "area" not in c_map:
                                c_map["area"] = c_i
                        elif "TOTAL GFA" in c_val:
                            c_map["perm_gfa"] = c_i
                            c_map["ach_gfa"] = c_i
                        elif any(k in c_val for k in ("PERMITTED GFA", "MAX GFA", "ALLOWABLE GFA", "PERMITTED BUA")):
                            c_map["perm_gfa"] = c_i
                        elif any(k in c_val for k in ("ACHIEVED GFA", "PROPOSED GFA", "GFA", "BUA")):
                            if "ach_gfa" not in c_map:
                                c_map["ach_gfa"] = c_i
                            if "perm_gfa" not in c_map:
                                c_map["perm_gfa"] = c_i
                        elif any(k in c_val for k in ("MAX FAR", "PERMITTED FAR", "MAXIMUM FAR")):
                            c_map["max_far"] = c_i
                        elif "FAR" in c_val:
                            c_map["far"] = c_i
                        elif any(k in c_val for k in ("HEIGHT", "FLOORS", "STORIES")):
                            c_map["height"] = c_i
                        elif "FRONT SETBACK" in c_val:
                            c_map["front_sb"] = c_i
                        elif "REAR SETBACK" in c_val:
                            c_map["rear_sb"] = c_i
                        elif "SIDE SETBACK" in c_val:
                            c_map["side_sb"] = c_i
                        elif "PARKING" in c_val:
                            c_map["parking"] = c_i
                    break

            if h_idx != -1 and "plot_id" in c_map:
                valid_count = 0
                for r in sh_rows[h_idx + 1:]:
                    if not r:
                        continue
                    pid_val = r[c_map["plot_id"]] if c_map["plot_id"] < len(r) else None
                    if not pid_val:
                        continue
                    s_pid = str(pid_val).strip().upper()
                    if any(ign in s_pid for ign in ("TOTAL", "SUBTOTAL", "AVERAGE", "SITEWIDE", "SUMMARY", "OVERALL", "RESIDENTIAL", "COMMERCIAL", "HOSPITALITY", "INDUSTRIAL", "RECREATIONAL", "UTILITIES")):
                        continue
                    valid_count += 1

                nu = name.upper()
                name_bonus = 100 if any(k in nu for k in ("AFTER", "PROPOSED", "SCHEDULE", "REGISTER", "PLOT")) else 0
                score = valid_count + name_bonus
                if score > best_score:
                    best_score = score
                    best_sheet = sh
                    best_header_idx = h_idx
                    best_col_map = c_map

        sheet = best_sheet
        header_idx = best_header_idx
        col_map = best_col_map
        rows = list(sheet.iter_rows(values_only=True))

        if not rows:
            return plots

        if header_idx == -1 or "plot_id" not in col_map:
            logger.warning(f"Header row not identified with high confidence in {sheet.title}; using fallback heuristic.")
            header_idx = 0
            col_map = {"plot_id": 0, "land_use": 1, "area": 2, "ach_gfa": 3, "far": 4}

        logger.info(f"Using sheet '{sheet.title}' for schedule data. Columns: {col_map} at header row {header_idx}")

        row_count = 0
        for row_number, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
            if not row or all(v is None for v in row):
                continue

            raw_id = row[col_map["plot_id"]] if "plot_id" in col_map and col_map["plot_id"] < len(row) else None
            if not raw_id:
                continue

            raw_id_str = str(raw_id).strip()
            raw_id_upper = raw_id_str.upper()

            # Ignore non-plot summary, category rows, and repeated table headers
            if any(k in raw_id_upper for k in (
                "TOTAL", "SUBTOTAL", "AVERAGE", "SITEWIDE", "SUMMARY", "OVERALL",
                "NOTE", "REMARK", "LEGEND", "RESIDENTIAL", "COMMERCIAL", "HOSPITALITY",
                "INDUSTRIAL", "RECREATIONAL", "UTILITIES", "MAXIMUM", "MINIMUM",
                "PLOT NO", "PLOT NUMBER", "PLOT ID"
            )) or raw_id_upper in ("PLOT NO.", "PLOT NO", "PLOT NUMBER", "PLOT #", "PLOT ID"):
                continue

            # Must have at least one alphanumeric character
            if not any(c.isalnum() for c in raw_id_str):
                continue

            plot_id_str = raw_id_str
            if not plot_id_str.upper().startswith("P-"):
                plot_id_str = f"P-{plot_id_str}"

            sector_str = str(row[col_map["sector"]]).strip() if "sector" in col_map and col_map["sector"] < len(row) and row[col_map["sector"]] else "Sector 1"
            land_use_str = str(row[col_map["land_use"]]).strip() if "land_use" in col_map and col_map["land_use"] < len(row) and row[col_map["land_use"]] else "Residential"

            area_val = _clean_num(row[col_map["area"]]) if "area" in col_map and col_map["area"] < len(row) else 1000.0
            perm_gfa = _clean_num(row[col_map["perm_gfa"]]) if "perm_gfa" in col_map and col_map["perm_gfa"] < len(row) else area_val * 2.0
            ach_gfa = _clean_num(row[col_map["ach_gfa"]]) if "ach_gfa" in col_map and col_map["ach_gfa"] < len(row) else perm_gfa

            max_far_val = _clean_num(row[col_map["max_far"]]) if "max_far" in col_map and col_map["max_far"] < len(row) else 2.5
            far_val = _clean_num(row[col_map["far"]]) if "far" in col_map and col_map["far"] < len(row) else (round(ach_gfa / area_val, 2) if area_val > 0 else 1.0)
            height_val = str(row[col_map["height"]]).strip() if "height" in col_map and col_map["height"] < len(row) and row[col_map["height"]] else "G+4"

            fs_val = _clean_num(row[col_map["front_sb"]], 6.0) if "front_sb" in col_map and col_map["front_sb"] < len(row) else 6.0
            rs_val = _clean_num(row[col_map["rear_sb"]], 4.5) if "rear_sb" in col_map and col_map["rear_sb"] < len(row) else 4.5
            ss_val = _clean_num(row[col_map["side_sb"]], 3.0) if "side_sb" in col_map and col_map["side_sb"] < len(row) else 3.0
            prk_val = int(_clean_num(row[col_map["parking"]], 0)) if "parking" in col_map and col_map["parking"] < len(row) else int(ach_gfa / 60)

            plot = MasterplanPlot(
                plot_id=plot_id_str,
                project_id=project_id,
                sector=sector_str,
                land_use=land_use_str,
                area_sqm=round(area_val, 1),
                permitted_gfa=round(perm_gfa, 1),
                achieved_gfa=round(ach_gfa, 1),
                far=round(far_val, 2),
                max_far=round(max_far_val, 2),
                max_height=height_val,
                max_coverage_pct=50.0,
                front_setback_m=fs_val,
                rear_setback_m=rs_val,
                side_setback_m=ss_val,
                parking_required=prk_val,
                parking_provided=prk_val,
                status="compliant",
                violations=[],
                provenance=[{"filename": os.path.basename(file_path), "sheet": sheet.title, "row": row_number, "kind": "excel_register"}],
                notes="Imported from Excel Plot Schedule"
            )
            plots.append(plot)
            row_count += 1

        wb.close()
        logger.info(f"MasterplanExcelParser extracted {len(plots)} plots from {file_path}")
        return plots
