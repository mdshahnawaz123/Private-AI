"""Structure-aware readers for Excel, Word and image files.

Same approach Claude's own xlsx/docx/pdf skills use: read the file as the
structure it really is (sheets with header rows, headings with sections,
tables with header cells, OCR for images) instead of flattening it into one
text blob. Each table row becomes a self-contained "[TABLE ROW]" unit that
carries its column headers, so retrieval and the answer checks always see the
value together with what it means.
"""
from __future__ import annotations

import datetime as _dt
import os
import re

MAX_ROW_UNITS_PER_TABLE = 3000


def _fmt(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, _dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    if isinstance(v, _dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return ("%.6f" % v).rstrip("0").rstrip(".")
    return re.sub(r"\s+", " ", str(v)).strip()


_NUM = re.compile(r"^[-+]?[\d,]*\.?\d+%?$")


def _is_texty(c):
    return bool(c) and not _NUM.match(c.replace(" ", ""))


def _is_data_row(row):
    filled = [c for c in row if c]
    if len(filled) < 2:
        return False
    nums = sum(1 for c in filled if not _is_texty(c))
    return nums / len(filled) >= 0.3


def _distinct(row):
    return len({c for c in row if c})


def _blocks(grid):
    """Split a sheet into blocks separated by empty rows."""
    out, cur = [], []
    for r in grid:
        if any(r):
            cur.append(r)
        elif cur:
            out.append(cur)
            cur = []
    if cur:
        out.append(cur)
    return out


def _headers_for(rows, width):
    """Combined column names from up to 3 header rows ("Group / Sub")."""
    heads = []
    for j in range(width):
        parts = []
        for r in rows:
            v = r[j] if j < len(r) else ""
            if v and v not in parts:
                parts.append(v)
        heads.append(" / ".join(parts) or "Column %d" % (j + 1))
    for j in range(1, len(heads)):
        if heads[j] == "Column %d" % (j + 1) and re.search(r"\b(ITEM|S\.?\s?NO\.?|NO\.?|#)$", heads[j - 1].strip(), re.I) \
                and "Description" not in heads:
            heads[j] = "Description"
    return heads


def _parse_block(block):
    """(title, headers, data_rows, kind). kind: 'table' | 'pairs' | 'text'."""
    width = max(len(r) for r in block)
    block = [list(r) + [""] * (width - len(r)) for r in block]
    titles = []
    while block and _distinct(block[0]) == 1 and not _is_data_row(block[0]):
        titles.append(next(c for c in block[0] if c))
        block = block[1:]
    title = " - ".join(titles)
    if not block:
        return title, [], [], "text"
    first_data = next((i for i, r in enumerate(block) if _is_data_row(r)), None)
    if first_data is None:
        return title, [], block, "text"
    head_rows = block[max(0, first_data - 3):first_data]
    if not head_rows:
        return title, [], block, "pairs"
    return title, _headers_for(head_rows, width), block[first_data:], "table"


def _is_total(row):
    first = next((c for c in row if c), "")
    return bool(re.match(r"(grand\s+)?(sub\s*)?total\b", first, re.I))


def table_row_units(grid, where, caption="", collect=None):
    """[TABLE ROW] units with headers attached, for every table block in the grid.
    Returns (units, summaries) -- summaries describe each block (columns + totals)."""
    units, summaries, pairs_found = [], [], []
    for block in _blocks(grid):
        title, heads, data, kind = _parse_block(block)
        label = " | ".join(x for x in (where, ("Table: " + (caption or title)) if (caption or title) else "") if x)
        if kind == "table":
            carry, k, pending = "", 0, []
            sections = [[label, heads, []]]
            for row in data:
                cells = list(row)
                filled = [c for c in cells if c]
                texty = [c for c in filled if _is_texty(c)]
                if not _is_data_row(cells):
                    # A sub-heading ("Proposed Residential") or a new header row inside the sheet.
                    if _distinct(cells) == 1 and not pending:
                        sec = filled[0]
                        label = " | ".join(x for x in (where, "Table: " + sec) if x)
                        sections.append([label, heads, []])
                        continue
                    if _distinct(cells) >= 2 and len(texty) >= 0.7 * len(filled) \
                            and not any(re.search(r"\d", c) for c in filled):
                        pending.append(cells)
                        continue
                if pending:
                    heads = _headers_for(pending[-3:], len(heads))
                    if sections[-1][2]:
                        sections.append([label, heads, []])
                    else:
                        sections[-1][1] = heads
                    pending = []
                k += 1
                tdist = {c for c in texty}
                if len(tdist) == 1 and len(texty) >= 2 and len(filled) - len(texty) <= 2:
                    # merged label cell + value(s): "PLOT AREA/ha.: 248.68"
                    vals = list(dict.fromkeys(c for c in filled if not _is_texty(c)))
                    pairs_found.append([label, texty[0], vals[0] if vals else ""])
                    units.append("[TABLE ROW] %s | Data row %d\n%s: %s" % (label, k, texty[0], "; ".join(vals)))
                    continue
                if not cells[0] and carry and any(cells[1:]):
                    cells[0] = carry
                carry = cells[0]
                pairs = ["%s: %s" % (heads[j], cells[j]) for j in range(len(heads)) if j < len(cells) and cells[j] and cells[j] != heads[j]]
                if pairs:
                    units.append("[TABLE ROW] %s | Data row %d\n%s" % (label, k, "; ".join(pairs)))
                    sections[-1][2].append(cells)
            for lab, hs, rows in sections:
                if not rows:
                    continue
                if collect is not None:
                    collect.append({"title": lab, "columns": list(hs), "rows": [list(r) for r in rows]})
                info = "%s - %d data rows. Columns: %s" % (lab, len(rows), "; ".join(dict.fromkeys(hs)))
                stats = _column_stats([r for r in rows if not _is_total(r)], hs, 0)
                if stats:
                    info += "\nColumn totals (computed, total rows excluded): " + " | ".join(stats)
                summaries.append(info)
        else:
            for k, row in enumerate(data, 1):
                vals = [c for c in row if c]
                if not vals:
                    continue
                if kind == "pairs" or (len(vals) >= 2 and _is_texty(vals[0])):
                    text = "%s: %s" % (vals[0], "; ".join(vals[1:])) if len(vals) > 1 else vals[0]
                    if len(vals) == 2:
                        pairs_found.append([label, vals[0], vals[1]])
                else:
                    text = " | ".join(vals)
                units.append("[TABLE ROW] %s | Data row %d\n%s" % (label, k, text))
        if len(units) >= MAX_ROW_UNITS_PER_TABLE:
            break
    if collect is not None and pairs_found:
        collect.append({"title": where, "pairs": pairs_found})
    return units, summaries


def _column_stats(rows, heads, start):
    """Count/sum/min/max for numeric columns -- the small model cannot add up 500 rows itself."""
    out = []
    for j, h in enumerate(heads):
        if "%" in h or re.search(r"\b(id|no\.?|number|plot|code|ref|year|date)\b", h, re.I):
            continue
        vals, filled = [], 0
        for r in rows[start:]:
            if j < len(r) and r[j]:
                filled += 1
                s = r[j].replace(",", "").replace(" ", "")
                if _NUM.match(s) and not s.endswith("%"):
                    try:
                        vals.append(float(s))
                    except ValueError:
                        pass
        if len(vals) >= 3 and len(vals) >= 0.6 * max(1, filled):
            out.append("%s: count %d, sum %s, min %s, max %s, average %s" % (
                h, len(vals), _fmt(round(sum(vals), 2)), _fmt(min(vals)), _fmt(max(vals)),
                _fmt(round(sum(vals) / len(vals), 2))))
    return out


# ---------------------------------------------------------------- Excel
def _cell_text(cell):
    v = cell.value
    fmt = str(getattr(cell, "number_format", "") or "")
    if isinstance(v, (int, float)) and not isinstance(v, bool) and "%" in fmt:
        return ("%.2f" % (v * 100)).rstrip("0").rstrip(".") + "%"
    if isinstance(v, float):
        return _fmt(round(v, 2) if abs(v) >= 100 else round(v, 4))
    return _fmt(v)


def _sheet_grids_openpyxl(path):
    import openpyxl
    big = os.path.getsize(path) > 25 * 1024 * 1024
    wb = openpyxl.load_workbook(path, data_only=True, read_only=big)
    try:
        for ws in wb.worksheets:
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            grid = [[_cell_text(c) for c in row] for row in ws.iter_rows()]
            if not big:
                for rng in ws.merged_cells.ranges:  # fill merged cells with the top-left value
                    v = grid[rng.min_row - 1][rng.min_col - 1] if rng.min_row - 1 < len(grid) and rng.min_col - 1 < len(grid[rng.min_row - 1]) else ""
                    for r in range(rng.min_row, rng.max_row + 1):
                        for c in range(rng.min_col, rng.max_col + 1):
                            if r - 1 < len(grid) and c - 1 < len(grid[r - 1]) and not grid[r - 1][c - 1]:
                                grid[r - 1][c - 1] = v
            yield ws.title, grid
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _sheet_grids_pandas(path):
    import pandas as pd
    for name, df in pd.read_excel(path, sheet_name=None, header=None).items():
        yield str(name), [[_fmt(None if (isinstance(v, float) and v != v) else v) for v in row]
                          for row in df.itertuples(index=False)]


def extract_excel(path):
    try:
        sheets_iter = list(_sheet_grids_openpyxl(path))
    except Exception:
        sheets_iter = list(_sheet_grids_pandas(path))  # .xls and odd files
    sheets, parts = [], []
    for title, grid in sheets_iter:
        if not any(any(r) for r in grid):
            continue
        width = max(len(r) for r in grid)
        grid = [list(r) + [""] * (width - len(r)) for r in grid]
        while width and not any(r[width - 1] for r in grid):  # trim empty right columns
            width -= 1
            grid = [r[:width] for r in grid]
        units, summaries = table_row_units(grid, "Sheet: " + title)
        summary = "[SHEET SUMMARY] Sheet: %s\n%s" % (title, "\n".join(summaries) or "no table header found")
        rows = [r for r in grid if any(r)]
        sheets.append({"sheet": title, "rows": rows, "summary": summary, "row_units": units,
                       "summaries": ["[SHEET SUMMARY] " + x for x in summaries]})
        parts.append("[SHEET %s]\n%s" % (title, "\n".join(" | ".join(c for c in r if c) for r in rows)))
    return "\n\n".join(parts), {"type": "excel", "sheets": sheets}


# ---------------------------------------------------------------- Word
def _docx_table_grid(tbl):
    grid = []
    for row in tbl.rows:
        cells, seen = [], set()
        for cell in row.cells:
            key = id(cell._tc)
            text = _fmt(cell.text)
            cells.append("" if key in seen else text)  # merged cell repeated by python-docx
            seen.add(key)
        grid.append(cells)
    return grid


def extract_docx(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    heading_path, sections, cur = [], [], {"heading": "", "lines": []}
    paras, tables, row_units, parts = [], [], [], []
    last_para = ""
    pending_caption_for = None
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = _fmt(p.text)
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            m = re.match(r"(?:Heading|Title)\s*(\d*)", style, re.I)
            if m:
                level = int(m.group(1) or 1)
                heading_path = heading_path[:level - 1] + [text]
                if cur["lines"]:
                    sections.append(cur)
                cur = {"heading": " > ".join(heading_path), "lines": []}
                paras.append("#" * min(level, 6) + " " + text)
                parts.append("#" * min(level, 6) + " " + text)
            else:
                if pending_caption_for is not None and re.match(r"(Table|Figure)\s*\d", text, re.I):
                    tables[pending_caption_for]["caption"] = text
                cur["lines"].append(text)
                paras.append(text)
                parts.append(text)
            pending_caption_for = None
            last_para = text
        elif tag == "tbl":
            grid = _docx_table_grid(Table(child, d))
            caption = last_para if re.match(r"(Table)\s*\d", last_para or "", re.I) else ""
            tables.append({"rows": grid, "caption": caption, "section": cur["heading"]})
            pending_caption_for = len(tables) - 1
            parts.append("\n".join(" | ".join(r) for r in grid))
    if cur["lines"]:
        sections.append(cur)
    for i, t in enumerate(tables):
        where = "Table %d" % (i + 1) + ((" | Section: " + t["section"]) if t["section"] else "")
        units, _ = table_row_units(t["rows"], where, t.get("caption", ""))
        row_units.extend(units)
    meta = {"type": "docx", "paragraphs": paras,
            "tables": [t["rows"] for t in tables], "table_info": tables,
            "sections": [{"heading": s["heading"], "text": "\n".join(s["lines"])} for s in sections],
            "row_units": row_units}
    return "\n\n".join(parts), meta


# ---------------------------------------------------------------- Images
def image_ocr(image_path, timeout=180):
    """OCR text (column-aware order) for a standalone image, cached beside it."""
    import json
    import subprocess
    import sys
    cache = image_path + ".ocr.json"
    if not os.path.isfile(cache):
        done = subprocess.run([sys.executable, "-m", "services.page_ocr", image_path, cache, "1"],
                              timeout=timeout, capture_output=True, text=True)
        if done.returncode:
            raise RuntimeError("Image OCR failed: " + done.stderr[-400:])
    with open(cache, encoding="utf-8") as handle:
        result = json.load(handle)
    text = result.get("text") or ""
    try:
        from services.ocr_layout import reading_order_text
        ordered = reading_order_text(result.get("boxes"))
        if ordered and ordered.strip():
            text = ordered
    except Exception:
        pass
    return text.strip()


def image_structured(image_path):
    """{'ocr_text', 'table_units'} for an image upload. Failures are non-fatal."""
    out = {"ocr_text": "", "table_units": []}
    try:
        out["ocr_text"] = image_ocr(image_path)
    except Exception as exc:
        out["ocr_error"] = str(exc)[:300]
    try:
        from services.image_tables import table_from_image
        grid = table_from_image(image_path)
        if grid:
            out["table_units"], _ = table_row_units(grid, "Image table")
    except Exception as exc:
        out["table_error"] = str(exc)[:300]
    return out
